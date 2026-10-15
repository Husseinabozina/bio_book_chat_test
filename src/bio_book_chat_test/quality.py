"""DOCX structure audit and optional render-based quality report."""
from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import subprocess
import shutil
import fitz
from docx import Document
from docx.oxml.ns import qn

from .analysis import inspect_pdf, suspicious_chars
from .ooxml_checks import docx_structure_warnings


def audit_docx(docx_file: str | Path, *, original_pdf: str | Path | None = None,
               pages: str | None = None, render: bool = False) -> dict:
    doc = Document(str(docx_file))
    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(p.text for p in paragraphs)
    rtl = sum(p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None for p in paragraphs)
    report = {
        "file": Path(docx_file).name,
        "editable_characters": len(text),
        "suspicious_unicode_in_docx": len(suspicious_chars(text)),
        "paragraphs_with_text": len(paragraphs),
        "paragraphs_marked_rtl": int(rtl),
        "tables": len(doc.tables),
        "inline_images": len(doc.inline_shapes),
        "warnings": docx_structure_warnings(docx_file),
    }
    if report["suspicious_unicode_in_docx"]:
        report["warnings"].append("OUTPUT_HAS_SUSPICIOUS_UNICODE")
    if not text:
        report["warnings"].append("NO_EDITABLE_TEXT_DETECTED")
    if paragraphs and rtl < max(1, len(paragraphs)//2):
        report["warnings"].append("MOST_PARAGRAPHS_NOT_RTL")
    if original_pdf:
        orig = inspect_pdf(original_pdf, pages)
        report["source_inspection"] = {
            "pages_inspected": orig["pages_inspected"],
            "requires_text_review": orig["requires_text_review"],
            "total_extracted_chars": sum(p["chars_extracted"] for p in orig["page_reports"]),
        }
    if render:
        if not shutil.which("libreoffice"):
            report["warnings"].append("RENDER_SKIPPED_LIBREOFFICE_UNAVAILABLE")
        else:
            with TemporaryDirectory() as tmp:
                cmd = ["libreoffice", "-env:UserInstallation=file://"+tmp+"/lo-profile", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(Path(docx_file).absolute())]
                p = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
                pdf = Path(tmp)/(Path(docx_file).stem+".pdf")
                if p.returncode or not pdf.exists():
                    report["warnings"].append("DOCX_RENDER_FAILED")
                    report["render_log"] = (p.stderr or p.stdout)[-1000:]
                else:
                    with fitz.open(str(pdf)) as converted:
                        report["rendered_page_count"] = len(converted)
                        if original_pdf and len(converted) != len(orig["pages_inspected"]):
                            report["warnings"].append("RENDERED_PAGE_COUNT_MISMATCH")
    report["status"] = "review_required" if report["warnings"] else "automated_checks_passed_manual_review_still_required"
    return report
