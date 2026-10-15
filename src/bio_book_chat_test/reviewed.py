"""Source-image-verified, editable Arabic DOCX from a declarative review manifest.

Unlike the PDF text baseline, this mode never trusts damaged PDF Unicode. A
reviewer supplies source-checked text and figure regions. Exported paragraphs,
headings, captions and sidebars are native Word objects; diagrams remain images.

Coordinates for figure crops are PDF *points*, not pixels from a scaled preview.
This small renderer is intentionally transparent rather than pixel-perfect.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import json
import re
import unicodedata

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Keep a multi-word Latin technical term (and its parentheses) on one line.
LATIN_TERM = re.compile(r"(\([A-Za-z0-9][A-Za-z0-9\s.,\-]*\)|[A-Za-z][A-Za-z0-9\s\-]*[A-Za-z0-9]|\d+(?:[.,]\d+)*)")


def validate_text(value: str) -> None:
    if not isinstance(value, str):
        raise TypeError("reviewed text must be a Unicode string")
    if "\ufffd" in value or any(0x80 <= ord(c) <= 0x9f for c in value):
        raise ValueError("unreviewed damaged Unicode: use source image to verify before export")
    if any(unicodedata.category(c) == "Cc" and c not in "\n\t" for c in value):
        raise ValueError("unexpected Unicode control in reviewed transcription")


def rtl(p, *, visual_heading: bool = False) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if not visual_heading:
        prop = p._p.get_or_add_pPr()
        el = OxmlElement("w:bidi");el.set(qn("w:val"), "1");prop.append(el)


def mixed_runs(p, value: str, *, font_size: float, bold: bool = False, color: RGBColor | None = None) -> None:
    validate_text(value)
    for piece in LATIN_TERM.split(value):
        if not piece:
            continue
        latin = bool(re.match(r"^\(?[A-Za-z0-9]", piece))
        if latin:
            piece = piece.replace(" ", "\N{NO-BREAK SPACE}")
        run = p.add_run(piece)
        run.bold = bold
        run.font.size = Pt(font_size)
        run.font.color.rgb = color or RGBColor(25, 25, 25)
        font = "Liberation Serif" if latin else "Amiri"
        run.font.name = font
        prop = run._r.get_or_add_rPr()
        fonts = prop.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts"); prop.insert(0, fonts)
        for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
            fonts.set(qn("w:" + attr), font)
        direction = OxmlElement("w:rtl")
        direction.set(qn("w:val"), "0" if latin else "1")
        prop.append(direction)


def add_text(container, content: str, *, kind: str = "paragraph", size: float | None = None):
    defaults = {"paragraph": 14.2, "caption": 12.0, "heading": 17.0, "sidebar": 12.5}
    if kind not in defaults:
        raise ValueError(f"Unsupported text kind: {kind}")
    p = container.add_paragraph()
    rtl(p, visual_heading=(kind == "heading"))
    p.paragraph_format.space_after = Pt(7 if kind == "paragraph" else 5)
    p.paragraph_format.line_spacing = 1.18
    if kind in ("paragraph", "sidebar"):
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if kind == "heading":
        sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), "E0EDF5")
        p._p.get_or_add_pPr().append(sh)
    if kind == "sidebar":
        sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), "E6F5FB")
        p._p.get_or_add_pPr().append(sh)
    mixed_runs(p, content, font_size=size or defaults[kind], bold=kind == "heading")
    return p


def figure_bytes(block: dict, base: Path, pdf: fitz.Document | None) -> bytes:
    if "file" in block:
        path = (base / block["file"]).resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        return path.read_bytes()
    if pdf is None:
        raise ValueError("figure using crop_pdf_pt requires input_pdf in manifest")
    pgnum = int(block["source_page"])
    if not 1 <= pgnum <= len(pdf):
        raise ValueError("invalid 1-based source_page")
    rect = fitz.Rect(block["crop_pdf_pt"])
    page = pdf[pgnum - 1]
    if rect.is_empty or not page.rect.contains(rect):
        raise ValueError("figure crop must be within source page and use PDF points")
    dpi = min(300, max(100, int(block.get("dpi", 200))))
    return page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), clip=rect, alpha=False).tobytes("png")


def add_figure(container, content: bytes, *, width_in: float = 6.3):
    p = container.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(BytesIO(content), width=Inches(min(6.85, max(0.5, width_in))))
    return p


def add_table(document: Document, block: dict) -> int:
    """Create a native editable Word table from explicit, source-verified cell strings."""
    rows = block.get("rows")
    if not isinstance(rows, list) or not rows or not isinstance(rows[0], list) or not rows[0]:
        raise ValueError("table.rows must be a nonempty list of nonempty rows")
    cols = len(rows[0])
    if any(not isinstance(row, list) or len(row) != cols for row in rows):
        raise ValueError("table rows must have equal length")
    for row in rows:
        for value in row:
            validate_text(value)
    tbl = document.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    direction = OxmlElement("w:bidiVisual")
    tbl._tbl.tblPr.append(direction)
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            p = cell.paragraphs[0]
            rtl(p)
            p.paragraph_format.space_after = Pt(3)
            mixed_runs(p, value, font_size=float(block.get("font_size", 11.5)), bold=bool(block.get("header", False) and i == 0))
    return sum(len(row) for row in rows)


def render_reviewed_manifest(manifest_path: str | Path, destination: str | Path) -> dict:
    manifest_path = Path(manifest_path).resolve()
    spec = json.loads(manifest_path.read_text(encoding="utf-8"))
    pages = spec.get("pages")
    if not isinstance(pages, list) or not pages:
        raise ValueError("manifest.pages must be a nonempty list")
    pdf = None
    if spec.get("input_pdf"):
        pdf = fitz.open(str((manifest_path.parent / spec["input_pdf"]).resolve()))
    try:
        document = Document()
        figure_count = 0
        text_count = 0
        table_count = 0
        table_cells = 0
        for i, page in enumerate(pages):
            if i:
                sec = document.add_section(WD_SECTION_START.NEW_PAGE)
            else:
                sec = document.sections[0]
            sec.page_height = Inches(11.69)
            sec.page_width = Inches(8.27)
            sec.left_margin = sec.right_margin = Inches(0.64)
            sec.top_margin = Inches(0.70)
            sec.bottom_margin = Inches(0.58)
            sec.footer_distance = Inches(0.36)
            sec.footer.is_linked_to_previous = False
            pgnum = int(page.get("source_page", i + 1))
            foot = sec.footer.paragraphs[0]
            rtl(foot, visual_heading=True)
            mixed_runs(foot, f"Original source page {pgnum}", font_size=9)
            for block in page.get("blocks", []):
                kind = block.get("kind")
                if kind in ("paragraph", "caption", "heading", "sidebar"):
                    add_text(document, block["text"], kind=kind, size=block.get("font_size"))
                    text_count += 1
                elif kind == "figure":
                    add_figure(document, figure_bytes(block, manifest_path.parent, pdf), width_in=float(block.get("width_in", 6.3)))
                    figure_count += 1
                elif kind == "table":
                    table_cells += add_table(document, block)
                    table_count += 1
                elif kind == "image_text_row":
                    table = document.add_table(rows=1, cols=2)
                    table.autofit = False
                    table.columns[0].width = Inches(2.25)
                    table.columns[1].width = Inches(4.60)
                    picture_cell, text_cell = table.rows[0].cells
                    picture_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    picture_cell.paragraphs[0].add_run().add_picture(BytesIO(figure_bytes(block, manifest_path.parent, pdf)), width=Inches(2.15))
                    p = text_cell.paragraphs[0]; rtl(p);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
                    mixed_runs(p, block["text"], font_size=float(block.get("font_size", 13.3)))
                    figure_count += 1; text_count += 1
                else:
                    raise ValueError(f"Unsupported block kind: {kind}")
        target = Path(destination)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)
        return {"pages_requested": len(pages), "editable_text_blocks": text_count, "native_tables": table_count, "native_table_cells": table_cells, "image_figures": figure_count,
                "output": str(target), "render_and_manual_review_required": True}
    finally:
        if pdf is not None:
            pdf.close()
