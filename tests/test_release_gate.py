"""Synthetic safety regressions. These tests DO NOT certify any biology-book output."""
from hashlib import sha256
import json
from pathlib import Path
from docx import Document

from bio_book_chat_test.release_gate import release_check
from bio_book_chat_test.cli import main


def fixture(tmp_path: Path):
    f = tmp_path / "sample.docx"
    d = Document()
    d.add_paragraph("تجربة (ATP) بسيطة")
    d.save(f)
    digest = sha256(f.read_bytes()).hexdigest()
    ref = {"status": "verified", "evidence": "synthetic-review-record", "reviewer": "test-reviewer"}
    data = {
        "schema_version": 1, "docx_sha256": digest,
        "source_pdf_sha256": "a" * 64,
        "reviews": {k: dict(ref) for k in ("source_text", "visual_layout", "illustrations", "editability")},
        "word_compatibility": {
            "application": "Microsoft Word", "version": "16.test",
            "result": "opened_edited_saved_without_repair",
            "evidence": "synthetic-only; not actual Word test", "tested_docx_sha256": digest},
        "pages": [{
            "source_pdf_page": 41,
            "background_provenance": "text_removed_artwork",
            "reading_order_reviewed": True,
            "bbox_tolerance_pt": 5,
            "objects": [
                {"id": "p41-b01", "kind": "paragraph", "paragraph_id": "p41-paragraph-1",
                 "text": "تجربة (ATP) بسيطة",
                 "runs": [
                     {"text": "تجربة (", "direction": "rtl"},
                     {"text": "ATP", "direction": "ltr"},
                     {"text": ") بسيطة", "direction": "rtl"}],
                 "source_bbox_pt": [10, 10, 200, 30],
                 "rendered_bbox_pt": [11, 10, 200, 31],
                 "text_review": {"status": "verified", "source_image_ref": "synthetic-page-41.png",
                                 "reviewer": "test-reviewer"},
                 "typography": {"calibration": "measured_source_and_render"}}]
        }]
    }
    manifest = tmp_path / "evidence.json"
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return f, manifest, data


def codes(result):
    return {x["code"] for x in result["issues"]}


def test_strict_gate_fails_without_actual_render(tmp_path):
    docx, manifest, _ = fixture(tmp_path)
    result = release_check(docx, manifest, render=False)
    assert result["decision"] == "BLOCKED"
    assert "RENDER_REQUIRED" in codes(result)
    assert result["certified_by_automation"] is False


def test_word_compatibility_requires_attested_exact_binary(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    data["word_compatibility"]["tested_docx_sha256"] = "b" * 64
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert "WORD_DESKTOP_NOT_VERIFIED" in codes(release_check(docx, manifest))


def test_fails_source_and_render_geometry_drift(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    data["pages"][0]["objects"][0]["rendered_bbox_pt"] = [100, 80, 300, 99]
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert "LAYOUT_DRIFT" in codes(release_check(docx, manifest))


def test_rejects_rtl_latin_mixed_in_one_run(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    data["pages"][0]["objects"][0]["runs"] = [
        {"text": "تجربة (ATP) بسيطة", "direction": "rtl"}]
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert "MIXED_SCRIPTS_IN_ONE_RUN" in codes(release_check(docx, manifest))
    assert "LATIN_DIRECTION_WRONG" in codes(release_check(docx, manifest))


def test_refuses_visual_only_document_even_with_claimed_reviews(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    Document().save(docx)
    data["docx_sha256"] = sha256(docx.read_bytes()).hexdigest()
    data["word_compatibility"]["tested_docx_sha256"] = data["docx_sha256"]
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    assert "NOT_EDITABLE" in codes(release_check(docx, manifest))
    assert "TEXT_NOT_IN_DOCX" in codes(release_check(docx, manifest))


def test_missing_review_and_line_level_anti_pattern_are_blocked(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    data["reviews"]["source_text"]["status"] = "not_checked"
    data["pages"][0]["objects"][0]["kind"] = "line"
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    found = codes(release_check(docx, manifest))
    assert "REVIEW_NOT_APPROVED" in found
    assert "SPLIT_PARAGRAPH_INTO_LINES" in found


def test_docx_modified_after_review_is_blocked(tmp_path):
    docx, manifest, _ = fixture(tmp_path)
    d = Document(docx); d.add_paragraph("extra"); d.save(docx)
    assert "ARTIFACT_HASH_MISMATCH" in codes(release_check(docx, manifest))


def test_cli_release_check_exits_nonzero_and_writes_report(tmp_path):
    docx, manifest, _ = fixture(tmp_path)
    report = tmp_path / "gate.json"
    assert main(["release-check", str(docx), str(manifest), "--report", str(report)]) == 3
    assert json.loads(report.read_text(encoding="utf-8"))["decision"] == "BLOCKED"


def test_undeclared_word_title_and_missing_font_measurements_block(tmp_path):
    docx, manifest, data = fixture(tmp_path)
    d = Document(docx)
    d.add_paragraph("عنوان من اختراع المولد")
    d.save(docx)
    data["docx_sha256"] = sha256(docx.read_bytes()).hexdigest()
    data["word_compatibility"]["tested_docx_sha256"] = data["docx_sha256"]
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    found = codes(release_check(docx, manifest))
    assert "WORD_HAS_UNDECLARED_TEXT" in found
    assert "FONT_METRICS_MISSING" in found


def test_synthetic_fully_attested_control_can_reach_human_approval(monkeypatch, tmp_path):
    # A synthetic control with a stubbed renderer: tests state machine only.
    # DOES NOT simulate or assert a real Microsoft Word review happened.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    import bio_book_chat_test.quality as qa
    docx, manifest, data = fixture(tmp_path)
    d = Document()
    p = d.add_paragraph()
    bidi = OxmlElement("w:bidi"); bidi.set(qn("w:val"), "1")
    p._p.get_or_add_pPr().append(bidi)
    for txt, direction in [("تجربة (", "1"), ("ATP", "0"), (") بسيطة", "1")]:
        run = p.add_run(txt)
        rprop = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), direction)
        rprop.append(rtl)
    d.save(docx)
    digest = sha256(docx.read_bytes()).hexdigest()
    data["docx_sha256"] = digest
    data["word_compatibility"]["tested_docx_sha256"] = digest
    data["pages"][0]["objects"][0]["typography"].update({
        "source_font": "Synthetic Arabic", "output_font": "Synthetic Arabic",
        "source_font_size_pt": 14.0, "render_font_size_pt": 14.0,
        "source_line_height_pt": 20.0, "render_line_height_pt": 20.0})
    manifest.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(qa, "audit_docx", lambda *a, **k: {"rendered_page_count": 1, "warnings": []})
    result = release_check(docx, manifest, render=True)
    assert result["decision"] == "ELIGIBLE_FOR_HUMAN_FINAL_APPROVAL", result["issues"]
    assert result["certified_by_automation"] is False
