"""Synthetic-only tests: do not place copied textbook content in public CI."""
import json
from pathlib import Path
import pytest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from bio_book_chat_test.reviewed import render_reviewed_manifest, validate_text


def test_rejects_broken_unicode():
    for value in ["حرف\ufffd", "علم\x81 أحياء"]:
        with pytest.raises(ValueError):
            validate_text(value)
    validate_text("امتصاص الماء (Xylem) في الخشب")


def test_two_page_source_reviewed_document(tmp_path: Path):
    Image.new("RGB", (200, 140), "white").save(tmp_path / "synthetic.png")
    manifest = {"pages": [
        {"source_page": 5, "blocks": [
            {"kind": "heading", "text": "1- عنوان علمي"},
            {"kind": "paragraph", "text": "نص تجريبي (Light reactions) مع اتجاه صحيح."},
            {"kind": "figure", "file": "synthetic.png", "width_in": 3.2}]},
        {"source_page": 7, "blocks": [
            {"kind": "image_text_row", "file": "synthetic.png", "text": "جدول مختلط (ATP, NADPH)"},
            {"kind": "caption", "text": "شكل (7-1) تجربة صناعية"}]}]}
    (tmp_path / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / "synthetic.docx"
    report = render_reviewed_manifest(tmp_path / "manifest.json", output)
    d = Document(output)
    assert report["pages_requested"] == 2
    assert report["image_figures"] == 2
    assert len(d.sections) == 2
    assert len(d.inline_shapes) == 2
    assert d.paragraphs[1]._p.pPr.find(qn("w:bidi")) is not None
    assert "نص تجريبي" in d.paragraphs[1].text
    assert "ATP" in d.tables[0].cell(0, 1).text


def test_native_table_is_editable_and_rtl(tmp_path: Path):
    manifest = {"pages": [{"source_page": 1, "blocks": [
        {"kind": "table", "header": True, "rows": [["المادة", "الرمز"], ["الماء", "H2O"], ["ثنائي أوكسيد الكاربون", "CO2"]]}
    ]}]}
    (tmp_path / "spec.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    report = render_reviewed_manifest(tmp_path / "spec.json", tmp_path / "out.docx")
    d = Document(tmp_path / "out.docx")
    assert report["native_tables"] == 1
    assert report["native_table_cells"] == 6
    assert d.tables[0].cell(1, 0).text == "الماء"
    assert d.tables[0].cell(2, 1).text == "CO2"
    assert d.tables[0]._tbl.tblPr.find(qn("w:bidiVisual")) is not None


def test_native_table_rejects_misaligned_rows(tmp_path: Path):
    manifest = {"pages": [{"blocks": [{"kind": "table", "rows": [["أ", "ب"], ["ج"]]}]}]}
    (tmp_path / "broken.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(ValueError, match="equal length"):
        render_reviewed_manifest(tmp_path / "broken.json", tmp_path / "out.docx")


def test_reviewed_cli(tmp_path: Path):
    from bio_book_chat_test.cli import main
    manifest = {"pages": [{"blocks": [{"kind": "paragraph", "text": "نص تجريبي عربي"}]}]}
    (tmp_path / "spec.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    code = main(["reviewed", str(tmp_path / "spec.json"), str(tmp_path / "out.docx"),
                 "--report", str(tmp_path / "report.json")])
    assert code == 0
    assert "نص تجريبي" in Document(tmp_path / "out.docx").paragraphs[0].text
    assert json.loads((tmp_path / "report.json").read_text(encoding="utf-8"))["pages_requested"] == 1


def test_no_invented_english_footer_and_rtl_heading(tmp_path: Path):
    manifest = {"pages": [{"source_page": 41, "blocks": [
        {"kind": "heading", "text": "عنوان عربي من العينة"},
        {"kind": "paragraph", "text": "فقرة صناعية"}]}]}
    (tmp_path / "m.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    target = tmp_path / "out.docx"
    result = render_reviewed_manifest(tmp_path / "m.json", target)
    doc = Document(target)
    assert result["release_status"] == "EXPERIMENTAL_NOT_APPROVED"
    assert "Original source page" not in "\n".join(p.text for sec in doc.sections for p in sec.footer.paragraphs)
    assert doc.paragraphs[0]._p.pPr.find(qn("w:bidi")) is not None
